"""
Suite de QA para Libros Registro Oficiales de la AEAT según Orden HAC/773/2019 (User Story 3).
Valida que se generen los libros oficiales en Excel con todas las columnas obligatorias.
"""

import io
import pytest
import openpyxl
from app.domain.services.tax_books_service import TaxBooksService


def test_qa_invoices_issued_official_book():
    """Valida la generación del Libro Registro de Facturas Expedidas según HAC/773/2019."""
    service = TaxBooksService()
    
    invoices = [
        {
            "series": "F2026",
            "number": 1,
            "issue_date": "2026-01-15",
            "operation_date": "2026-01-15",
            "recipient_nif": "B12345674",
            "recipient_name": "ACME S.L.",
            "invoice_type": "F1",
            "base_amount": 1000.0,
            "vat_rate": 21.0,
            "tax_amount": 210.0,
            "retention_amount": 0.0,
            "total_amount": 1210.0
        },
        {
            "series": "F2026",
            "number": 2,
            "issue_date": "2026-02-10",
            "operation_date": "2026-02-10",
            "recipient_nif": "12345678Z",
            "recipient_name": "JUAN PEREZ",
            "invoice_type": "F1",
            "base_amount": 500.0,
            "vat_rate": 10.0,
            "tax_amount": 50.0,
            "retention_amount": 75.0,  # 15% retención
            "total_amount": 475.0
        }
    ]
    
    excel_bytes = service.generate_invoices_issued_book(2026, invoices)
    assert len(excel_bytes) > 0
    
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet = wb.active
    
    # Validar cabeceras obligatorias según Orden HAC/773/2019
    header_values = [cell.value for cell in sheet[1]]
    assert "Fecha Expedición" in header_values or "Fecha" in header_values
    assert "Número Factura" in header_values or "Número" in header_values
    assert "NIF Destinatario" in header_values or "NIF" in header_values
    assert "Base Imponible" in header_values
    assert "Tipo IVA" in header_values
    assert "Cuota IVA" in header_values
    assert "Total Factura" in header_values
    
    # Validar filas de datos
    assert sheet.max_row >= 3
    assert sheet.cell(row=2, column=header_values.index("Total Factura") + 1).value == 1210.0


def test_qa_invoices_received_official_book():
    """Valida la generación del Libro Registro de Facturas Recibidas según HAC/773/2019."""
    service = TaxBooksService()
    
    expenses = [
        {
            "reception_number": 1,
            "invoice_number": "PRV-9988",
            "issue_date": "2026-01-20",
            "supplier_nif": "A98765432",
            "supplier_name": "PROVEEDOR TECNOLOGICO S.A.",
            "base_amount": 800.0,
            "vat_rate": 21.0,
            "tax_amount": 168.0,
            "deductible_tax_amount": 168.0,
            "total_amount": 968.0
        }
    ]
    
    excel_bytes = service.generate_invoices_received_book(2026, expenses)
    assert len(excel_bytes) > 0
    
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet = wb.active
    header_values = [cell.value for cell in sheet[1]]
    
    assert "NIF Proveedor" in header_values or "NIF" in header_values
    assert "Cuota Deducible" in header_values or "IVA Deducible" in header_values


def test_qa_investment_goods_and_income_expense_books():
    """Valida la generación del Libro de Bienes de Inversión y el Libro de Ingresos y Gastos."""
    service = TaxBooksService()
    
    assets = [
        {
            "code": "ACT-001",
            "description": "Servidor IA Local",
            "acquisition_date": "2026-01-01",
            "acquisition_value": 3000.0,
            "depreciation_rate": 25.0,
            "annual_depreciation": 750.0
        }
    ]
    
    # Bienes de Inversión
    assets_bytes = service.generate_investment_goods_book(2026, assets)
    assert len(assets_bytes) > 0
    wb_assets = openpyxl.load_workbook(io.BytesIO(assets_bytes))
    assert wb_assets.active.max_row >= 2
    
    # Ingresos y Gastos
    inc_exp_bytes = service.generate_expenses_and_income_book(2026, incomes=[], expenses=[])
    assert len(inc_exp_bytes) > 0
    wb_inc = openpyxl.load_workbook(io.BytesIO(inc_exp_bytes))
    assert wb_inc.active.max_row >= 1
